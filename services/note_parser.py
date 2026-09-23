"""Normalizes uploaded notes (PDF, DOCX, PPTX, TXT) into plain text."""
import io
from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def parse_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def parse_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def parse_pptx(file_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(file_bytes))
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                chunks.append(shape.text_frame.text)
            if shape.has_notes_frame if hasattr(shape, "has_notes_frame") else False:
                chunks.append(shape.notes_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            note_text = slide.notes_slide.notes_text_frame.text
            if note_text.strip():
                chunks.append(note_text)
    return "\n".join(chunks)


def parse_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="ignore")


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".pptx": parse_pptx,
    ".txt": parse_txt,
}


def parse_uploaded_file(filename: str, file_bytes: bytes) -> str:
    ext = "." + filename.rsplit(".", 1)[-1].lower()
    parser = PARSERS.get(ext)
    if not parser:
        raise ValueError(f"Unsupported file type: {ext}. Supported: {list(PARSERS)}")
    text = parser(file_bytes)
    if not text.strip():
        raise ValueError("No extractable text found in this file.")
    return text
